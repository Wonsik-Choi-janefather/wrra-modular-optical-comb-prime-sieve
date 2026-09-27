#!/usr/bin/env python3
"""Build the Korean DOCX and PDF research reports."""

from __future__ import annotations

import os
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
FIGURES = ROOT / "figures"
DOCX_OUT = ROOT / "report" / "WRRA_모듈러_광주파수빗_소수체_구현연구_최원식_2026-09-27.docx"
PDF_OUT = ROOT / "report" / "WRRA_모듈러_광주파수빗_소수체_구현연구_최원식_2026-09-27.pdf"

FONT = os.environ.get("WRRA_FONT_REGULAR", "/usr/local/share/fonts/wrra/NotoSansKR-Regular.ttf")
FONT_BOLD = os.environ.get("WRRA_FONT_BOLD", "/usr/local/share/fonts/wrra/NotoSansKR-Bold.ttf")
DOCX_FONT = os.environ.get("WRRA_DOCX_FONT", "Noto Sans CJK KR")
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = "595959"


TITLE = "WRRA 기반 모듈러 광주파수 빗 소수체 구현 연구"
SUBTITLE = "잔여류 선택 필터의 정확성 제어 복잡도 및 오차 민감도 계산"
AUTHOR = "최원식 Wonsik Choi"
DATE = "2026년 9월 27일"


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    new_run.append(rpr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    new_run.append(text_element)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def style_docx(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOCX_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, before, after in [
        ("Title", 22, 0, 9),
        ("Subtitle", 12.5, 0, 15),
        ("Heading 1", 15.5, 14, 6),
        ("Heading 2", 12, 9, 4),
        ("Heading 3", 11, 8, 4),
    ]:
        st = styles[name]
        st.font.name = DOCX_FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
        st.font.size = Pt(size)
        st.font.bold = name != "Subtitle"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = f"{AUTHOR}   |   {DATE}   |   "
    for run in p.runs:
        run.font.name = DOCX_FONT
        run._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(p)


def add_docx_para(doc: Document, text: str, bold_lead: str | None = None, align=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead) :])
    else:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    return p


def add_docx_bullets(doc: Document, items: list[str]):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)
        p.paragraph_format.space_after = Pt(2)


def add_docx_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = header
        shade_cell(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.5)
        set_cell_margins(cell)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            cells[j].text = str(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 1:
                shade_cell(cells[j], PALE_BLUE)
            for p in cells[j].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 or len(str(value)) < 14 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.3)
            set_cell_margins(cells[j])
    if widths:
        for row in table.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_docx_figure(doc: Document, path: Path, caption: str, width: float = 6.8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for r in cap.runs:
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(80, 80, 80)


def add_docx_equation(doc: Document, filename: str, width: float):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(FIGURES / "equations" / filename), width=Inches(width))


def build_docx(complexity: pd.DataFrame, noise: pd.DataFrame):
    doc = Document()
    style_docx(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(TITLE)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(SUBTITLE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(AUTHOR).bold = True
    p.add_run(f"\n{DATE}")
    p.paragraph_format.space_after = Pt(22)

    doc.add_heading("연구 결론", level=1)
    add_docx_para(
        doc,
        "본 연구는 소수체의 배수 제거를 개별 comb line 제어가 아니라 잔여류 선택 필터의 연쇄로 표현할 수 있음을 확인했다. "
        "이상적 전달함수에서는 N이 30, 210, 2310일 때 모두 소수 마스크가 수치오차 수준에서 정확히 복원됐다. "
        "N 2310에서는 2309개의 독립 선 가중치 대신 15개의 주기 필터와 15개의 보호선 복원 동작으로 같은 이산 마스크를 기술했다. "
        "다만 이 결과는 산술 항등식과 축소 전달함수 시뮬레이션에 한정된다. 실제 광주파수 빗 장비에서의 구현 가능성은 아직 검증되지 않았다.",
    )
    add_docx_table(
        doc,
        ["판정 대상", "현재 판정"],
        [
            ["산술적 정확성", "증명 및 수치 검증 완료"],
            ["축소 제어 표현", "모델 내부에서 성립"],
            ["잡음 민감도", "축소 잡음모델에서 계산 완료"],
            ["실제 광학 구현", "미검증 가설"],
            ["수론상의 새 정리", "주장하지 않음"],
        ],
        widths=[1.7, 4.9],
    )

    doc.add_heading("초록", level=1)
    add_docx_para(
        doc,
        "기존 Prime-Sieve optical comb 제안은 산술 상태를 comb line에 대응시키고 선별적인 진폭 및 위상 제어로 체 과정을 구현하려 했다. "
        "이 방식의 직접적인 난점은 line-by-line 제어 부담이다. 본 연구는 WRRA Core의 Worldline, Residue, Resource, Action 구성을 이용해 "
        "각 소수 p의 배수선을 한 번에 제거하는 이산 푸리에 잔여류 선택 필터를 구성했다. p 이하의 p개 지연 탭 평균은 p의 배수에서만 1이므로, "
        "이를 1에서 빼면 p의 배수에 영점을 갖는 정확한 전달함수가 된다. N 이하의 모든 합성수는 제곱근 N 이하의 소인수를 가지므로 이 필터들을 "
        "연쇄하고 작은 소수선을 복원하면 이상적 조건에서 소수 지시함수를 정확히 얻는다. 30, 210, 2310 상태 계산에서 상대 L2 오차는 각각 "
        "3.91 x 10^-15, 3.84 x 10^-14, 4.50 x 10^-13이었다. 제어 기술은 크게 줄었지만, 지연오차가 선 번호에 따라 누적되면서 큰 상태공간에서 먼저 성능을 제한했다. "
        "따라서 다음 단계는 실제 장비 요청이 아니라 30선 규모의 수동 또는 프로그램형 지연 필터 설계와 전체 손실 예산 계산이다.",
    )

    doc.add_page_break()
    doc.add_heading("1 연구 배경", level=1)
    add_docx_para(
        doc,
        "Prime-Sieve State-Space Renormalization and Irreducibility 연구는 소수 체의 상태공간 재구성과 Fourier 관측량을 물리적 testbed에 매핑했다. "
        "광학 장치는 수론 정리를 증명하기 위한 장치가 아니라, 이미 정의된 이산 규칙이 물리적 전달함수에서 재현되는지를 반증 가능하게 시험하는 수단이다. "
        "장윤수 박사의 검토는 이 구성이 물리적으로 배제될 이유는 없지만 실제 frequency comb 제어가 어렵다는 점을 지적했다. "
        "본 연구는 이 지적을 새로운 입력 조건으로 사용한다.",
    )
    add_docx_para(
        doc,
        "WRRA 적용의 목적은 기존 광학 공학을 다른 이름으로 부르는 데 있지 않다. 동일한 관측 결과를 만드는 데 필요한 상태와 제어를 줄이고, "
        "제어 뒤에 남는 오차를 residue로 기록하며, 실패 조건을 수치로 드러내는 것이 목적이다.",
    )

    doc.add_heading("2 연구 질문과 가설", level=1)
    add_docx_bullets(
        doc,
        [
            "정확성 질문  주기적인 잔여류 필터의 곱으로 N 이하의 소수 마스크를 정확히 만들 수 있는가",
            "복잡도 질문  개별 comb line 설정 수에 비해 필요한 고수준 제어 동작과 광학 경로 수를 줄일 수 있는가",
            "강건성 질문  계수오차 위상오차 지연오차가 커질 때 소수와 합성수의 출력 대비가 어디에서 무너지는가",
            "반증 조건  동일한 제어 예산의 단순 묶음 또는 개별 억제보다 낫지 않거나 현실적인 오차에서 대비가 사라지면 적용 가설은 기각한다",
        ],
    )

    doc.add_heading("3 WRRA 도메인 구성", level=1)
    add_docx_table(
        doc,
        ["WRRA 요소", "광주파수 빗 대응", "계산에서의 역할"],
        [
            ["Worldline", "필터 단계별 comb line 상태", "각 p 필터 전후의 상태 이력"],
            ["Residue", "남은 합성수 출력과 손상된 소수 출력", "누설전력과 prime power 손실"],
            ["Resource", "필터 수 지연 탭 위상 정밀도", "제어 예산과 오차 범위"],
            ["Action", "p 주기 필터 적용과 보호선 복원", "배수 제거와 작은 소수 복원"],
            ["Renderer", "복소 전달함수의 연쇄", "최종 진폭 위상 및 Fourier 출력"],
            ["Ledger", "단계별 신규 제거와 잔여 합성수", "중복 제거와 실패 추적"],
        ],
        widths=[1.0, 2.5, 3.2],
    )

    doc.add_page_break()
    doc.add_heading("4 정확한 잔여류 선택 필터", level=1)
    add_docx_para(
        doc,
        "정수 n을 comb line 번호에 대응시킨다. 소수 p에 대한 p탭 이산 푸리에 평균을 이용해 다음 전달함수를 정의한다.",
    )
    add_docx_equation(doc, "equation_1_filter.png", 4.9)
    add_docx_para(
        doc,
        "근의 합 항등식에 의해 평균은 n이 p의 배수일 때만 1이다. 따라서 H는 p의 배수에서 0이고 다른 잔여류에서는 1이다.",
    )
    add_docx_equation(doc, "equation_2_selector.png", 5.6)
    add_docx_para(
        doc,
        "N 이하의 합성수는 반드시 제곱근 N 이하의 소인수를 갖는다. 그러므로 모든 p 이하 제곱근 N에 대한 필터를 곱하면 모든 합성수가 제거된다. "
        "각 필터는 p 자체도 제거하므로 작은 소수선을 별도의 보호 채널로 복원한다.",
    )
    add_docx_equation(doc, "equation_3_mask.png", 6.6)
    add_docx_para(
        doc,
        "이 식은 이산 comb grid 위에서 정확하다. 연속 주파수 전체에서 직사각형 마스크를 만드는 식이 아니며, 실제 광학장치가 이상적인 복소 탭 계수를 구현할 수 있다는 뜻도 아니다.",
    )

    doc.add_heading("5 제어 복잡도", level=1)
    add_docx_equation(doc, "equation_5_scaling.png", 5.5)
    add_docx_para(
        doc,
        "고수준 동작 수는 필터 적용과 작은 소수 복원을 합해 2 pi sqrt N으로 계산했다. 광학 경로 대용치는 각 p 필터의 직접 경로와 p개 탭, "
        "그리고 보호선 복원 채널을 합산했다. 이 값은 실제 칩 면적이나 삽입손실이 아니라 첫 설계를 비교하기 위한 구조적 계수다.",
    )
    rows = []
    for _, r in complexity.iterrows():
        rows.append([
            str(int(r.N)),
            str(int(r.prime_count)),
            str(int(r.direct_line_weights)),
            str(int(r.wrra_high_level_actions)),
            str(int(r.wrra_optical_paths_surrogate)),
            f"{r.action_reduction_vs_line_weights_pct:.1f}%",
        ])
    add_docx_table(
        doc,
        ["N", "소수 수", "독립 선 설정", "WRRA 동작", "경로 대용치", "동작 감소"],
        rows,
        widths=[0.6, 0.8, 1.2, 1.0, 1.2, 1.0],
    )
    add_docx_figure(doc, FIGURES / "figure_1_control_complexity.png", "그림 1 제어 및 구조 복잡도 비교", width=6.7)

    doc.add_page_break()
    doc.add_heading("6 계산 방법", level=1)
    add_docx_para(
        doc,
        "검증 규모는 기존 연구 계보와 연결되는 30, 210, 2310 상태로 정했다. 각 규모에서 이상적 출력, 동일한 고수준 동작 예산으로 일부 합성수만 개별 제거하는 기준선, "
        "동일 개수의 균등 구간에 상수 가중치를 주는 기준선을 비교했다. 이상적 필터는 결정론적으로 계산했다.",
    )
    add_docx_para(
        doc,
        "잡음 계산은 난수 시드 20260927을 고정했다. 시행 수는 N 30에서 400회, N 210에서 200회, N 2310에서 40회다. "
        "첫 번째 잡음군은 각 탭의 상대 계수오차와 고정 위상오차의 표준편차를 같은 수치로 두었다. 두 번째 잡음군은 지연의 상대오차를 독립적으로 적용했고, "
        "그 위상오차가 comb line 번호 n에 비례해 누적되도록 했다.",
    )
    add_docx_equation(doc, "equation_4_noisy.png", 6.5)
    add_docx_para(
        doc,
        "측정량은 최적 전역 복소 스케일 정렬 뒤의 상대 L2 오차, Fourier power correlation, 평균 합성수 누설, 소수 대 합성수 대비, "
        "소수선 출력의 5백분위수, 분류 precision recall F1이다. 보수적 통과 기준은 평균 F1과 5백분위 F1이 모두 0.95 이상이고 5백분위 대비가 20 dB 이상인 경우다.",
    )

    doc.add_heading("7 계산 결과", level=1)
    doc.add_heading("7 1 이상적 정확성", level=2)
    ideal_rows = []
    for _, r in complexity.iterrows():
        ideal_rows.append([
            str(int(r.N)),
            f"{r.ideal_relative_l2_error:.3e}",
            f"{r.ideal_fourier_correlation:.12f}",
            f"{r.wrra_ideal_f1:.3f}",
        ])
    add_docx_table(doc, ["N", "상대 L2 오차", "Fourier 상관", "F1"], ideal_rows, widths=[0.8, 1.5, 1.7, 0.9])
    add_docx_figure(doc, FIGURES / "figure_2_ideal_exactness.png", "그림 2 N 210에서 이상적 소수 마스크와 모듈러 필터 출력", width=6.8)

    doc.add_heading("7 2 동일 제어 예산 기준선", level=2)
    baseline_rows = []
    for _, r in complexity.iterrows():
        baseline_rows.append([
            str(int(r.N)),
            str(int(r.wrra_high_level_actions)),
            f"{r.budget_individual_f1:.3f}",
            f"{r.uniform_group_f1:.3f}",
            f"{r.wrra_ideal_f1:.3f}",
        ])
    add_docx_table(
        doc,
        ["N", "동작 예산", "개별 억제 F1", "균등 묶음 F1", "WRRA F1"],
        baseline_rows,
        widths=[0.7, 1.0, 1.3, 1.3, 1.0],
    )
    add_docx_para(
        doc,
        "WRRA 방식의 이득은 단순히 제어 수를 줄인 결과가 아니다. p의 배수라는 산술 구조를 하나의 물리 동작으로 소유하게 했기 때문에 같은 고수준 동작 예산에서 정확한 마스크가 가능했다. "
        "다만 각 주기 필터 내부의 탭 수를 포함하면 하드웨어 감소폭은 고수준 동작 감소폭보다 작다.",
    )

    doc.add_page_break()
    doc.add_heading("7 3 계수 및 위상오차", level=2)
    add_docx_figure(doc, FIGURES / "figure_3_coefficient_phase_robustness.png", "그림 3 탭 계수 및 고정 위상오차에 대한 강건성", width=6.8)
    add_docx_para(
        doc,
        "세 규모 모두 오차 표준편차 0.1까지 평균 F1은 거의 1을 유지했다. 그러나 0.3에서는 N 2310의 평균 F1이 0.896, 평균 대비가 17.16 dB로 떨어졌다. "
        "보수적 기준에서 확인된 최대 시험 통과치는 N 30이 0.03, N 210과 N 2310이 0.1이었다. 이 수치는 동일한 계수 표준편차와 라디안 위상 표준편차를 묶은 "
        "본 모델의 무차원 값이며 실제 장비 사양으로 직접 옮길 수 없다.",
    )

    doc.add_heading("7 4 지연오차", level=2)
    add_docx_figure(doc, FIGURES / "figure_4_delay_robustness.png", "그림 4 지연오차 누적에 대한 강건성", width=6.8)
    add_docx_para(
        doc,
        "지연오차는 상태공간이 커질수록 빠르게 누적됐다. 보수적 기준의 최대 시험 통과치는 N 30에서 1.0 x 10^-3, N 210에서 3.0 x 10^-4, "
        "N 2310에서 3.0 x 10^-5였다. N 2310에서 지연 상대오차 1.0 x 10^-4를 적용하면 평균 F1은 0.911, 평균 대비는 19.76 dB로 내려가 통과 기준을 벗어났다. "
        "따라서 실제 설계에서는 탭 계수 보정보다 경로 길이와 군지연 정렬을 먼저 다뤄야 한다.",
    )

    doc.add_heading("7 5 단계별 장부", level=2)
    ledger = pd.read_csv(RESULTS / "ideal_stage_ledger.csv").query("N == 210")
    ledger_rows = [[
        str(int(r.prime_filter_p)),
        str(int(r.newly_removed_composites)),
        str(int(r.remaining_composites)),
        str(int(r.protected_prime_temporarily_removed)),
    ] for _, r in ledger.iterrows()]
    add_docx_table(
        doc,
        ["필터 p", "새 합성수 제거", "남은 합성수", "임시 제거 소수"],
        ledger_rows,
        widths=[0.9, 1.4, 1.4, 1.4],
    )
    add_docx_para(
        doc,
        "N 210에서는 p 2 필터가 합성수 104개를 처음 제거하고, 이후 3 5 7 11 13 필터가 각각 34 13 7 4 1개를 새로 제거했다. "
        "마지막 단계 뒤 남은 합성수는 0이다. 이미 제거된 배수에 대한 후속 필터의 작용은 새 제거로 계산하지 않아 residue와 중복 동작을 분리했다.",
    )

    doc.add_page_break()
    doc.add_heading("8 결과 해석", level=1)
    add_docx_para(
        doc,
        "확인된 핵심은 소수 마스크가 개별 선 주소 목록으로만 표현될 필요가 없다는 점이다. 배수 구조를 주기 전달함수로 소유하면 적은 수의 고수준 동작으로 동일한 이산 출력을 기술할 수 있다. "
        "이 결과는 WRRA의 최소계산 직관과 일치한다. 복잡한 최종 표현형을 모두 저장하는 대신, 생성 규칙과 현재 단계의 residue를 유지한다.",
    )
    add_docx_para(
        doc,
        "그러나 이것이 WRRA가 기존 광학 최적화보다 우수하다는 증거는 아니다. 제안한 필터는 이산 푸리에 잔여류 선택기의 직접적인 응용이며, 유사한 tapped-delay-line 및 interferometric filter 선행기술이 존재한다. "
        "현재 새로 확인한 부분은 Prime-Sieve 상태공간과 WRRA 장부에 맞춘 정확한 조합, 제어계수 비교, 오류경계 계산이다. 광학적 신규성과 특허 가능성은 별도의 선행기술 조사가 필요하다.",
    )

    doc.add_heading("9 실험으로 옮기기 위한 최소 단계", level=1)
    add_docx_bullets(
        doc,
        [
            "N 30 파일럿부터 시작해 p 2 3 5 필터와 세 개의 보호 소수선만 구현한다",
            "실제 소자의 탭 계수 위상 군지연 삽입손실 열적 간섭을 측정해 본 잡음모델의 변수로 다시 넣는다",
            "소수 마스크만 측정하지 않고 무작위 마스크와 같은 밀도의 비소수 대조군을 함께 측정한다",
            "line power와 Fourier power를 모두 기록하고 이론에 맞지 않는 결과도 ledger에 남긴다",
            "N 30에서 20 dB 대비와 F1 0.95를 넘기기 전에는 N 210으로 확장하지 않는다",
        ],
    )
    add_docx_para(
        doc,
        "장윤수 박사에게 다시 문의할 수 있는 시점은 이 N 30 설계의 실제 지연허용오차, 삽입손실, 필요한 광원과 검출기 범위를 한 장으로 제시할 수 있을 때다. "
        "그때의 질문은 장비 사용 요청이 아니라 계산한 경로 길이 및 위상 정밀도가 실제 계측 환경에서 가능한지에 한정해야 한다.",
    )

    doc.add_page_break()
    doc.add_heading("10 한계", level=1)
    add_docx_bullets(
        doc,
        [
            "레이저 발진과 comb 생성 동역학을 포함하지 않았다",
            "비선형 전파 검출기 대역폭 피드백 전자회로를 포함하지 않았다",
            "삽입손실 열적 간섭 편광 의존성과 제작공차를 포함하지 않았다",
            "작은 소수선을 복원하는 보호 채널은 이상적으로 두었고 그 오류와 손실은 계산하지 않았다",
            "모든 탭의 오차를 독립 정규분포로 두었으므로 실제 상관오차와 장기 drift를 반영하지 못했다",
            "실험장치의 비용 면적 소비전력과 안정화 시간은 평가하지 않았다",
            "광학 구현은 소수 분포에 관한 새로운 정리나 리만가설의 증명을 제공하지 않는다",
        ],
    )

    doc.add_heading("11 주장 등급", level=1)
    add_docx_table(
        doc,
        ["주장", "등급", "근거"],
        [
            ["잔여류 필터 항등식", "Exact", "근의 합 항등식"],
            ["N 이하 소수 마스크 복원", "Exact", "합성수의 작은 소인수 존재"],
            ["30 210 2310 수치 재현", "Verified computation", "공개 코드와 고정 시드"],
            ["제어계수 감소", "Derived in surrogate", "정의한 동작 및 경로 계수"],
            ["오류 허용범위", "Model conditional", "단순화한 Monte Carlo 모델"],
            ["실제 광학장치 구현", "Open", "장비 설계와 실험 없음"],
            ["기존 방식 대비 공학적 우위", "Open", "동일 장비 기준 비교 없음"],
        ],
        widths=[2.1, 1.4, 3.0],
    )

    doc.add_page_break()
    doc.add_heading("12 결론", level=1)
    add_docx_para(
        doc,
        "이번 계산은 Prime-Sieve optical comb 아이디어를 개별 선 제어 문제에서 주기 필터 합성 문제로 바꾸었다. 산술 층에서는 정확한 구성이 완성됐고, 축소 모델에서는 제어 표현의 감소도 확인됐다. "
        "동시에 N 2310에서 지연오차가 먼저 한계를 만든다는 구체적인 실패점도 얻었다. 따라서 다음 연구의 우선순위는 새로운 우주론 해석이나 추가 수론 주장이 아니라 N 30 모듈의 물리적 전달함수와 손실 예산이다.",
    )

    doc.add_heading("참고문헌", level=1)
    refs = [
        "1  Choi W  Prime-Sieve State-Space Renormalization and Irreducibility  Projective Lifting Information Increment and Falsifiable Physical Mappings  Zenodo 2026  DOI 10.5281/zenodo.21869614",
        "2  Jiang Z  Huang C B  Leaird D E  Weiner A M  Line-by-line pulse shaping control for optical arbitrary waveform generation  Optics Express 13 10431-10439 2005  DOI 10.1364/OPEX.13.010431",
        "3  Cohen L M et al  Silicon photonic microresonator-based high-resolution line-by-line pulse shaping  Nature Communications 15 7878 2024  DOI 10.1038/s41467-024-52051-9",
        "4  Hong S et al  Versatile parallel signal processing with a scalable silicon photonic chip  Nature Communications 16 288 2025  DOI 10.1038/s41467-024-55162-5",
        "5  Ziyadi M et al  Tunable radio frequency photonics filter using a comb-based optical tapped delay line with an optical nonlinear multiplexer  Optics Letters 40 3284-3287 2015  DOI 10.1364/OL.40.003284",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)

    doc.add_heading("재현 정보", level=1)
    add_docx_para(
        doc,
        "계산 코드는 wrra_modular_comb.py에 수록했다. 전체 시행값, 집계표, 단계별 ledger, 그림 생성 코드와 난수 시드는 별도 재현 패키지에 포함했다. "
        "주요 결과를 재현하려면 Python 환경에서 해당 스크립트를 실행하면 된다.",
    )
    p = doc.add_paragraph()
    p.add_run("관련 선행 연구 주소  ").bold = True
    add_hyperlink(p, "https://doi.org/10.5281/zenodo.21869614", "https://doi.org/10.5281/zenodo.21869614")

    # LibreOffice may not consistently inherit East Asian font and bold settings
    # from heading styles, so make the run-level instruction explicit.
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            for run in paragraph.runs:
                run.font.name = DOCX_FONT
                run._element.rPr.rFonts.set(qn("w:eastAsia"), DOCX_FONT)
                run.font.bold = True

    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)


def pdf_styles():
    pdfmetrics.registerFont(TTFont("Malgun", FONT))
    pdfmetrics.registerFont(TTFont("MalgunBold", FONT_BOLD))
    return {
        "title": ParagraphStyle("TitleK", fontName="MalgunBold", fontSize=21, leading=28, alignment=TA_CENTER, textColor=colors.black, spaceAfter=10),
        "subtitle": ParagraphStyle("SubtitleK", fontName="Malgun", fontSize=12, leading=17, alignment=TA_CENTER, textColor=colors.black, spaceAfter=18),
        "author": ParagraphStyle("AuthorK", fontName="Malgun", fontSize=10, leading=15, alignment=TA_CENTER, textColor=colors.black, spaceAfter=20),
        "h1": ParagraphStyle("H1K", fontName="MalgunBold", fontSize=15, leading=20, textColor=colors.black, spaceBefore=12, spaceAfter=7, keepWithNext=True),
        "h2": ParagraphStyle("H2K", fontName="MalgunBold", fontSize=11.5, leading=16, textColor=colors.black, spaceBefore=9, spaceAfter=5, keepWithNext=True),
        "body": ParagraphStyle("BodyK", fontName="Malgun", fontSize=9.5, leading=15, alignment=TA_JUSTIFY, textColor=colors.black, spaceAfter=6),
        "small": ParagraphStyle("SmallK", fontName="Malgun", fontSize=8, leading=11, alignment=TA_LEFT, textColor=colors.HexColor("#505050"), spaceAfter=5),
        "table_header": ParagraphStyle("TableHeaderK", fontName="MalgunBold", fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.white),
        "bullet": ParagraphStyle("BulletK", fontName="Malgun", fontSize=9.4, leading=14, leftIndent=13, firstLineIndent=-8, spaceAfter=3),
    }


def pdf_table(headers, rows, widths, styles):
    data = [[Paragraph(str(h), styles["table_header"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(v), styles["small"]) for v in row])
    table = Table(data, colWidths=[w * inch for w in widths], repeatRows=1, hAlign="CENTER")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365D")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9D9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for i in range(1, len(data)):
        if i % 2 == 0:
            commands.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#EAF2F8")))
    table.setStyle(TableStyle(commands))
    return table


def build_pdf(complexity: pd.DataFrame):
    styles = pdf_styles()
    doc = BaseDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.65 * inch,
        title=TITLE,
        author=AUTHOR,
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")

    def on_page(canvas, document):
        canvas.saveState()
        canvas.setFont("Malgun", 7.5)
        canvas.setFillColor(colors.HexColor("#666666"))
        canvas.drawString(doc.leftMargin, 0.37 * inch, f"{AUTHOR}   {DATE}")
        canvas.drawRightString(letter[0] - doc.rightMargin, 0.37 * inch, str(document.page))
        canvas.restoreState()

    doc.addPageTemplates(PageTemplate(id="report", frames=frame, onPage=on_page))
    story = []

    def h1(text): story.append(Paragraph(text, styles["h1"]))
    def h2(text): story.append(Paragraph(text, styles["h2"]))
    def body(text): story.append(Paragraph(text, styles["body"]))
    def bullet(text): story.append(Paragraph(f"• {text}", styles["bullet"]))
    def fig(path, caption, width=6.6):
        im = Image(str(path), width=width * inch, height=width * inch * 0.58)
        im.hAlign = "CENTER"
        story.extend([im, Paragraph(caption, styles["small"]), Spacer(1, 5)])
    def eq(name, width, height):
        im = Image(str(FIGURES / "equations" / name), width=width * inch, height=height * inch)
        im.hAlign = "CENTER"
        story.extend([im, Spacer(1, 4)])

    story.extend([
        Spacer(1, 0.35 * inch),
        Paragraph(TITLE, styles["title"]),
        Paragraph(SUBTITLE, styles["subtitle"]),
        Paragraph(f"{AUTHOR}<br/>{DATE}", styles["author"]),
    ])
    h1("연구 결론")
    body("본 연구는 소수체의 배수 제거를 개별 comb line 제어가 아니라 잔여류 선택 필터의 연쇄로 표현할 수 있음을 확인했다. 이상적 전달함수에서는 N이 30, 210, 2310일 때 모두 소수 마스크가 수치오차 수준에서 정확히 복원됐다. N 2310에서는 2309개의 독립 선 가중치 대신 15개의 주기 필터와 15개의 보호선 복원 동작으로 같은 이산 마스크를 기술했다. 다만 이 결과는 산술 항등식과 축소 전달함수 시뮬레이션에 한정되며 실제 장비 구현은 미검증이다.")
    story.append(pdf_table(["판정 대상", "현재 판정"], [
        ["산술적 정확성", "증명 및 수치 검증 완료"], ["축소 제어 표현", "모델 내부에서 성립"],
        ["잡음 민감도", "축소 잡음모델에서 계산 완료"], ["실제 광학 구현", "미검증 가설"],
        ["수론상의 새 정리", "주장하지 않음"],
    ], [1.65, 5.15], styles))
    h1("초록")
    body("기존 Prime-Sieve optical comb 제안의 직접적인 난점은 line-by-line 제어 부담이다. 본 연구는 WRRA Core의 Worldline, Residue, Resource, Action 구성을 이용해 각 소수 p의 배수선을 한 번에 제거하는 이산 푸리에 잔여류 선택 필터를 구성했다. p개의 지연 탭 평균은 p의 배수에서만 1이므로 이를 1에서 빼면 배수에 영점을 갖는 전달함수가 된다. N 이하의 모든 합성수는 제곱근 N 이하의 소인수를 가지므로 필터들을 연쇄하고 작은 소수선을 복원하면 이상적 조건에서 소수 지시함수를 정확히 얻는다. 30, 210, 2310 상태의 상대 L2 오차는 각각 3.91 x 10^-15, 3.84 x 10^-14, 4.50 x 10^-13이었다. 제어 기술은 줄었지만 지연오차가 선 번호에 따라 누적되면서 큰 상태공간에서 먼저 성능을 제한했다.")

    story.append(PageBreak())
    h1("1 연구 배경")
    body("Prime-Sieve State-Space Renormalization and Irreducibility 연구는 소수 체의 상태공간 재구성과 Fourier 관측량을 물리적 testbed에 매핑했다. 광학 장치는 수론을 증명하는 장치가 아니라 이미 정의된 이산 규칙이 물리적 전달함수에서 재현되는지를 시험하는 수단이다. 장윤수 박사의 검토는 이 구성이 물리적으로 배제될 이유는 없지만 실제 frequency comb 제어가 어렵다는 점을 지적했다. 본 연구는 이 지적을 새로운 입력 조건으로 사용한다.")
    body("WRRA 적용의 목적은 동일한 관측 결과를 만드는 데 필요한 상태와 제어를 줄이고, 제어 뒤에 남는 오차를 residue로 기록하며, 실패 조건을 수치로 드러내는 것이다.")
    h1("2 연구 질문과 가설")
    for item in [
        "주기적인 잔여류 필터의 곱으로 N 이하의 소수 마스크를 정확히 만들 수 있는가",
        "개별 comb line 설정 수에 비해 필요한 고수준 제어 동작과 광학 경로 수를 줄일 수 있는가",
        "계수오차 위상오차 지연오차가 커질 때 소수와 합성수의 출력 대비가 어디에서 무너지는가",
        "동일 제어 예산의 단순 묶음 또는 개별 억제보다 낫지 않으면 적용 가설을 기각한다",
    ]: bullet(item)
    h1("3 WRRA 도메인 구성")
    story.append(pdf_table(["WRRA 요소", "광주파수 빗 대응", "계산 역할"], [
        ["Worldline", "단계별 comb line 상태", "필터 전후 상태 이력"],
        ["Residue", "남은 합성수와 손상된 소수", "누설과 prime power 손실"],
        ["Resource", "필터 탭 위상 정밀도", "제어 예산과 오차"],
        ["Action", "주기 필터와 보호선 복원", "배수 제거와 소수 복원"],
        ["Renderer", "복소 전달함수 연쇄", "최종 출력과 Fourier 관측"],
        ["Ledger", "신규 제거와 잔여 합성수", "중복 및 실패 추적"],
    ], [1.0, 2.6, 3.2], styles))

    story.append(PageBreak())
    h1("4 정확한 잔여류 선택 필터")
    body("정수 n을 comb line 번호에 대응시킨다. 소수 p에 대한 p탭 이산 푸리에 평균을 이용해 다음 전달함수를 정의한다.")
    eq("equation_1_filter.png", 5.0, 0.45)
    body("근의 합 항등식에 의해 평균은 n이 p의 배수일 때만 1이다. 따라서 H는 p의 배수에서 0이고 다른 잔여류에서는 1이다.")
    eq("equation_2_selector.png", 5.5, 0.47)
    body("N 이하의 합성수는 반드시 제곱근 N 이하의 소인수를 갖는다. 모든 해당 필터를 곱하면 합성수가 제거되며, 필터가 함께 제거한 작은 소수선은 보호 채널로 복원한다.")
    eq("equation_3_mask.png", 6.6, 0.52)
    body("이 식은 이산 comb grid 위에서 정확하다. 실제 광학장치가 이상적인 복소 탭 계수를 구현할 수 있다는 뜻은 아니다.")
    h1("5 제어 복잡도")
    eq("equation_5_scaling.png", 5.6, 0.42)
    body("고수준 동작 수는 필터 적용과 작은 소수 복원을 합해 계산했다. 광학 경로 대용치는 각 p 필터의 직접 경로와 p개 탭, 보호선 복원 채널을 합산했다. 이 값은 실제 칩 면적이나 삽입손실이 아니다.")
    rows = [[str(int(r.N)), str(int(r.prime_count)), str(int(r.direct_line_weights)), str(int(r.wrra_high_level_actions)), str(int(r.wrra_optical_paths_surrogate)), f"{r.action_reduction_vs_line_weights_pct:.1f}%"] for _, r in complexity.iterrows()]
    story.append(pdf_table(["N", "소수 수", "독립 선 설정", "WRRA 동작", "경로 대용치", "동작 감소"], rows, [0.55, 0.75, 1.25, 1.05, 1.2, 1.0], styles))
    fig(FIGURES / "figure_1_control_complexity.png", "그림 1 제어 및 구조 복잡도 비교")

    story.append(PageBreak())
    h1("6 계산 방법")
    body("검증 규모는 30, 210, 2310 상태로 정했다. 이상적 출력, 같은 고수준 동작 예산으로 일부 합성수만 개별 제거하는 기준선, 동일 개수의 균등 구간에 상수 가중치를 주는 기준선을 비교했다. 잡음 계산은 시드 20260927을 고정하고 각 탭의 계수오차 고정 위상오차 지연 상대오차를 적용했다.")
    eq("equation_4_noisy.png", 6.6, 0.48)
    body("측정량은 상대 L2 오차, Fourier power correlation, 합성수 누설, 소수 대 합성수 대비, prime power 5백분위수, precision recall F1이다. 보수적 통과 기준은 평균 F1과 5백분위 F1이 0.95 이상이고 5백분위 대비가 20 dB 이상인 경우다.")
    h1("7 계산 결과")
    h2("7 1 이상적 정확성")
    rows = [[str(int(r.N)), f"{r.ideal_relative_l2_error:.3e}", f"{r.ideal_fourier_correlation:.12f}", f"{r.wrra_ideal_f1:.3f}"] for _, r in complexity.iterrows()]
    story.append(pdf_table(["N", "상대 L2 오차", "Fourier 상관", "F1"], rows, [0.8, 1.6, 1.7, 0.8], styles))
    fig(FIGURES / "figure_2_ideal_exactness.png", "그림 2 N 210에서 이상적 소수 마스크와 모듈러 필터 출력", 6.7)
    h2("7 2 동일 제어 예산 기준선")
    rows = [[str(int(r.N)), str(int(r.wrra_high_level_actions)), f"{r.budget_individual_f1:.3f}", f"{r.uniform_group_f1:.3f}", f"{r.wrra_ideal_f1:.3f}"] for _, r in complexity.iterrows()]
    story.append(pdf_table(["N", "동작 예산", "개별 억제 F1", "균등 묶음 F1", "WRRA F1"], rows, [0.7, 1.0, 1.3, 1.3, 0.9], styles))
    body("WRRA 방식은 p의 배수라는 산술 구조를 하나의 물리 동작으로 소유해 같은 고수준 동작 예산에서 정확한 마스크를 만들었다. 각 주기 필터의 내부 탭 수를 포함하면 하드웨어 감소폭은 고수준 동작 감소폭보다 작다.")

    story.append(PageBreak())
    h2("7 3 계수 및 위상오차")
    fig(FIGURES / "figure_3_coefficient_phase_robustness.png", "그림 3 탭 계수 및 고정 위상오차에 대한 강건성", 6.7)
    body("세 규모 모두 오차 표준편차 0.1까지 평균 F1은 거의 1을 유지했다. 0.3에서는 N 2310의 평균 F1이 0.896, 평균 대비가 17.16 dB로 떨어졌다. 보수적 최대 시험 통과치는 N 30이 0.03, N 210과 N 2310이 0.1이었다. 이 수치는 본 무차원 모델 안에서만 해석해야 한다.")
    h2("7 4 지연오차")
    fig(FIGURES / "figure_4_delay_robustness.png", "그림 4 지연오차 누적에 대한 강건성", 6.7)
    body("지연오차는 상태공간이 커질수록 빠르게 누적됐다. 보수적 최대 시험 통과치는 N 30에서 1.0 x 10^-3, N 210에서 3.0 x 10^-4, N 2310에서 3.0 x 10^-5였다. N 2310에서 1.0 x 10^-4를 적용하면 평균 F1은 0.911, 평균 대비는 19.76 dB로 내려갔다.")
    h2("7 5 단계별 장부")
    ledger = pd.read_csv(RESULTS / "ideal_stage_ledger.csv").query("N == 210")
    rows = [[str(int(r.prime_filter_p)), str(int(r.newly_removed_composites)), str(int(r.remaining_composites)), str(int(r.protected_prime_temporarily_removed))] for _, r in ledger.iterrows()]
    story.append(pdf_table(["필터 p", "새 합성수 제거", "남은 합성수", "임시 제거 소수"], rows, [0.85, 1.45, 1.4, 1.4], styles))

    h1("8 결과 해석")
    body("소수 마스크는 개별 선 주소 목록으로만 표현될 필요가 없다. 배수 구조를 주기 전달함수로 소유하면 적은 수의 고수준 동작으로 같은 이산 출력을 기술할 수 있다. 복잡한 최종 표현형을 모두 저장하지 않고 생성 규칙과 현재 residue를 유지한다는 점에서 WRRA의 최소계산 직관과 일치한다.")
    body("그러나 이것이 기존 광학 최적화보다 우수하다는 증거는 아니다. 이산 푸리에 잔여류 선택기와 tapped-delay-line 및 interferometric filter 선행기술이 존재한다. 현재 확인한 부분은 Prime-Sieve 상태공간과 WRRA 장부에 맞춘 정확한 조합, 제어계수 비교, 오류경계 계산이다. 광학적 신규성은 별도 선행기술 조사가 필요하다.")
    h1("9 실험으로 옮기기 위한 최소 단계")
    for item in [
        "N 30 파일럿에서 p 2 3 5 필터와 세 보호 소수선만 구현한다",
        "탭 계수 위상 군지연 삽입손실 열적 간섭을 측정해 모델 변수로 다시 넣는다",
        "소수 마스크와 무작위 대조 마스크를 함께 측정한다",
        "line power와 Fourier power를 함께 기록하고 실패 결과도 ledger에 남긴다",
        "N 30에서 20 dB 대비와 F1 0.95를 넘기기 전에는 N 210으로 확장하지 않는다",
    ]: bullet(item)
    body("장윤수 박사에게 다시 문의할 수 있는 시점은 N 30 설계의 지연허용오차, 삽입손실, 필요한 광원과 검출기 범위를 한 장으로 제시할 수 있을 때다. 질문은 장비 사용 요청이 아니라 계산한 경로 길이와 위상 정밀도가 실제 계측 환경에서 가능한지에 한정한다.")

    h1("10 한계")
    for item in [
        "레이저 발진과 comb 생성 동역학을 포함하지 않았다",
        "비선형 전파 검출기 대역폭 피드백 전자회로를 포함하지 않았다",
        "삽입손실 열적 간섭 편광 의존성과 제작공차를 포함하지 않았다",
        "보호선 복원 채널의 오류와 손실을 계산하지 않았다",
        "독립 정규오차를 사용해 실제 상관오차와 장기 drift를 반영하지 못했다",
        "광학 구현은 소수 분포의 새 정리나 리만가설 증명을 제공하지 않는다",
    ]: bullet(item)
    h1("11 주장 등급")
    story.append(pdf_table(["주장", "등급", "근거"], [
        ["잔여류 필터 항등식", "Exact", "근의 합 항등식"],
        ["N 이하 소수 마스크 복원", "Exact", "작은 소인수 존재"],
        ["30 210 2310 수치 재현", "Verified computation", "코드와 고정 시드"],
        ["제어계수 감소", "Derived in surrogate", "정의한 구조 계수"],
        ["오류 허용범위", "Model conditional", "Monte Carlo 모델"],
        ["실제 광학 구현", "Open", "장비 실험 없음"],
        ["공학적 우위", "Open", "동일 장비 비교 없음"],
    ], [2.1, 1.45, 3.1], styles))
    h1("12 결론")
    body("이번 계산은 Prime-Sieve optical comb 아이디어를 개별 선 제어 문제에서 주기 필터 합성 문제로 바꾸었다. 산술 층에서는 정확한 구성이 완성됐고 축소 모델에서는 제어 표현의 감소도 확인됐다. 동시에 N 2310에서 지연오차가 먼저 한계를 만든다는 실패점도 얻었다. 다음 우선순위는 N 30 모듈의 물리적 전달함수와 손실 예산이다.")
    h1("참고문헌")
    refs = [
        "1  Choi W  Prime-Sieve State-Space Renormalization and Irreducibility  Zenodo 2026  DOI 10.5281/zenodo.21869614",
        "2  Jiang Z et al  Line-by-line pulse shaping control for optical arbitrary waveform generation  Optics Express 13 10431-10439 2005  DOI 10.1364/OPEX.13.010431",
        "3  Cohen L M et al  Silicon photonic microresonator-based high-resolution line-by-line pulse shaping  Nature Communications 15 7878 2024  DOI 10.1038/s41467-024-52051-9",
        "4  Hong S et al  Versatile parallel signal processing with a scalable silicon photonic chip  Nature Communications 16 288 2025  DOI 10.1038/s41467-024-55162-5",
        "5  Ziyadi M et al  Tunable radio frequency photonics filter using a comb-based optical tapped delay line  Optics Letters 40 3284-3287 2015  DOI 10.1364/OL.40.003284",
    ]
    for ref in refs: story.append(Paragraph(ref, styles["small"]))
    h1("재현 정보")
    body("계산 코드, 전체 시행값, 집계표, 단계별 ledger, 그림 생성 코드와 난수 시드는 별도 재현 패키지에 포함했다. 주요 결과는 Python에서 wrra_modular_comb.py를 실행해 재현할 수 있다.")

    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)


def main():
    complexity = pd.read_csv(RESULTS / "complexity_and_ideal_results.csv")
    noise = pd.read_csv(RESULTS / "noise_monte_carlo_summary.csv")
    build_docx(complexity, noise)
    build_pdf(complexity)
    print(DOCX_OUT)
    print(PDF_OUT)


if __name__ == "__main__":
    main()
